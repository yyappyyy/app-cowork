#!/usr/bin/env python3
"""Generate a professional household budget spreadsheet (家計簿)."""

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.chart import PieChart, BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from copy import copy

OUTPUT_PATH = "/home/runner/work/app-cowork/app-cowork/household-budget-spreadsheet/家計簿テンプレート.xlsx"

# Color scheme
COLORS = {
    "primary": "2C3E50",
    "accent": "3498DB",
    "success": "27AE60",
    "warning": "F39C12",
    "danger": "E74C3C",
    "light_bg": "ECF0F1",
    "white": "FFFFFF",
    "header_bg": "34495E",
    "income_bg": "D5F5E3",
    "expense_bg": "FADBD8",
    "savings_bg": "D6EAF8",
}

# Styles
header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
title_font = Font(name="Arial", size=16, bold=True, color=COLORS["primary"])
subtitle_font = Font(name="Arial", size=11, bold=True, color=COLORS["primary"])
normal_font = Font(name="Arial", size=10)
header_fill = PatternFill(start_color=COLORS["header_bg"], end_color=COLORS["header_bg"], fill_type="solid")
income_fill = PatternFill(start_color=COLORS["income_bg"], end_color=COLORS["income_bg"], fill_type="solid")
expense_fill = PatternFill(start_color=COLORS["expense_bg"], end_color=COLORS["expense_bg"], fill_type="solid")
savings_fill = PatternFill(start_color=COLORS["savings_bg"], end_color=COLORS["savings_bg"], fill_type="solid")
light_fill = PatternFill(start_color=COLORS["light_bg"], end_color=COLORS["light_bg"], fill_type="solid")

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

MONTHS = ["1月", "2月", "3月", "4月", "5月", "6月", "7月", "8月", "9月", "10月", "11月", "12月"]
MONTH_ICONS = ["⛄", "💝", "🌸", "🌷", "🎏", "☔", "🎋", "🌻", "🌾", "🎃", "🍂", "🎄"]

INCOME_CATEGORIES = ["💼 給与", "🔨 副業", "📈 投資収入", "🎁 その他収入"]
EXPENSE_CATEGORIES = [
    "🏠 住居費（家賃/ローン）",
    "💡 水道光熱費",
    "🍽️ 食費",
    "🧴 日用品",
    "🚃 交通費",
    "📱 通信費",
    "🛡️ 保険料",
    "🏥 医療費",
    "📚 教育費",
    "🎮 娯楽費",
    "👔 衣服費",
    "💇 美容費",
    "🍻 交際費",
    "📦 その他",
]


def style_header_row(ws, row, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def create_dashboard(wb):
    ws = wb.active
    ws.title = "ダッシュボード"
    set_col_widths(ws, [2, 18, 16, 16, 16, 16, 16, 16, 2])

    # Row heights for better spacing
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 8

    # ===== Title =====
    ws.merge_cells("B2:H2")
    ws["B2"] = "📊 家計簿ダッシュボード"
    ws["B2"].font = Font(name="Arial", size=20, bold=True, color=COLORS["primary"])
    ws["B2"].alignment = Alignment(vertical="center")

    # ===== KPI Cards Row =====
    # 4 KPI cards: 年間収入, 年間支出, 年間貯蓄, 貯蓄率
    kpi_row = 4
    ws.row_dimensions[kpi_row].height = 18
    ws.row_dimensions[kpi_row + 1].height = 30
    ws.row_dimensions[kpi_row + 2].height = 14

    kpi_configs = [
        ("💰 年間収入", "income_bg", "B", "C"),
        ("💸 年間支出", "expense_bg", "D", "E"),
        ("🏦 年間貯蓄", "savings_bg", "F", "G"),
        ("📈 貯蓄率", "light_bg", "H", "H"),
    ]

    for i, (label, bg_key, col_start, col_end) in enumerate(kpi_configs):
        col_num = 2 + i * 2 if i < 3 else 8
        # Label row
        cell = ws.cell(row=kpi_row, column=col_num, value=label)
        cell.font = Font(name="Arial", size=9, color="666666")
        cell.fill = PatternFill(start_color=COLORS[bg_key], end_color=COLORS[bg_key], fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        if i < 3:
            ws.cell(row=kpi_row, column=col_num + 1).fill = PatternFill(
                start_color=COLORS[bg_key], end_color=COLORS[bg_key], fill_type="solid")

    # KPI value row
    val_row = kpi_row + 1
    # Income total
    ws.merge_cells(f"B{val_row}:C{val_row}")
    ws.cell(row=val_row, column=2, value=f"=SUM(C{kpi_row+6}:C{kpi_row+17})")
    ws.cell(row=val_row, column=2).font = Font(name="Arial", size=16, bold=True, color=COLORS["success"])
    ws.cell(row=val_row, column=2).number_format = '#,##0"円"'
    ws.cell(row=val_row, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=val_row, column=2).fill = income_fill
    ws.cell(row=val_row, column=3).fill = income_fill

    # Expense total
    ws.merge_cells(f"D{val_row}:E{val_row}")
    ws.cell(row=val_row, column=4, value=f"=SUM(D{kpi_row+6}:D{kpi_row+17})")
    ws.cell(row=val_row, column=4).font = Font(name="Arial", size=16, bold=True, color=COLORS["danger"])
    ws.cell(row=val_row, column=4).number_format = '#,##0"円"'
    ws.cell(row=val_row, column=4).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=val_row, column=4).fill = expense_fill
    ws.cell(row=val_row, column=5).fill = expense_fill

    # Savings total
    ws.merge_cells(f"F{val_row}:G{val_row}")
    ws.cell(row=val_row, column=6, value=f"=B{val_row}-D{val_row}")
    ws.cell(row=val_row, column=6).font = Font(name="Arial", size=16, bold=True, color=COLORS["accent"])
    ws.cell(row=val_row, column=6).number_format = '#,##0"円"'
    ws.cell(row=val_row, column=6).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=val_row, column=6).fill = savings_fill
    ws.cell(row=val_row, column=7).fill = savings_fill

    # Savings rate
    ws.cell(row=val_row, column=8, value=f'=IF(B{val_row}=0,"--",F{val_row}/B{val_row})')
    ws.cell(row=val_row, column=8).font = Font(name="Arial", size=16, bold=True, color=COLORS["primary"])
    ws.cell(row=val_row, column=8).number_format = "0%"
    ws.cell(row=val_row, column=8).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=val_row, column=8).fill = light_fill

    # Add borders to KPI cards
    for r in range(kpi_row, val_row + 1):
        for c in range(2, 9):
            ws.cell(row=r, column=c).border = thin_border

    # ===== Separator =====
    sep_row = kpi_row + 3
    ws.row_dimensions[sep_row].height = 8

    # ===== Section: Monthly Summary Table =====
    section_row = kpi_row + 4
    ws.merge_cells(f"B{section_row}:H{section_row}")
    ws[f"B{section_row}"] = "📅 月別収支サマリー"
    ws[f"B{section_row}"].font = Font(name="Arial", size=13, bold=True, color=COLORS["primary"])
    ws.row_dimensions[section_row].height = 22

    # Table headers
    tbl_hdr_row = section_row + 1
    headers2 = ["🗓️ 月", "💚 収入", "🔴 支出", "💰 貯蓄", "📊 貯蓄率", "📉 前月比支出", "🏅 評価"]
    for col, h in enumerate(headers2, 2):
        ws.cell(row=tbl_hdr_row, column=col, value=h)
    style_header_row(ws, tbl_hdr_row, 8)

    # Monthly data rows
    for i, month in enumerate(MONTHS):
        r = tbl_hdr_row + 1 + i
        ws.cell(row=r, column=2, value=f"{MONTH_ICONS[i]} {month}").border = thin_border
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2).font = Font(name="Arial", size=10, bold=True)
        # Reference formulas to monthly sheets
        income_ref = f"'{month}'!E3"
        expense_ref = f"'{month}'!E10"
        ws.cell(row=r, column=3).value = f"={income_ref}"
        ws.cell(row=r, column=3).number_format = '#,##0"円"'
        ws.cell(row=r, column=4).value = f"={expense_ref}"
        ws.cell(row=r, column=4).number_format = '#,##0"円"'
        ws.cell(row=r, column=5).value = f"={income_ref}-{expense_ref}"
        ws.cell(row=r, column=5).number_format = '#,##0"円"'
        ws.cell(row=r, column=6).value = f'=IF({income_ref}=0,"",({income_ref}-{expense_ref})/{income_ref})'
        ws.cell(row=r, column=6).number_format = "0%"
        # Month-over-month expense comparison
        if i == 0:
            ws.cell(row=r, column=7).value = '="--"'
        else:
            ws.cell(row=r, column=7).value = f'=IF(D{r-1}=0,"--",(D{r}-D{r-1})/D{r-1})'
            ws.cell(row=r, column=7).number_format = "+0%;-0%"
        # Evaluation emoji
        ws.cell(row=r, column=8).value = f'=IF({income_ref}=0,"",IF(({income_ref}-{expense_ref})/{income_ref}>=0.2,"◎ 優秀",IF(({income_ref}-{expense_ref})/{income_ref}>=0.1,"○ 良好",IF(({income_ref}-{expense_ref})>=0,"△ 普通","✕ 赤字"))))'
        ws.cell(row=r, column=8).alignment = Alignment(horizontal="center")
        for c in range(2, 9):
            ws.cell(row=r, column=c).border = thin_border
        # Alternate row shading
        if i % 2 == 0:
            for c in range(2, 9):
                ws.cell(row=r, column=c).fill = PatternFill(
                    start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

    # ===== Charts Section =====
    chart_start_row = tbl_hdr_row + 14

    # Bar Chart - Monthly income vs expense
    ws.merge_cells(f"B{chart_start_row}:D{chart_start_row}")
    ws[f"B{chart_start_row}"] = "📊 月別 収入 vs 支出 推移"
    ws[f"B{chart_start_row}"].font = Font(name="Arial", size=12, bold=True, color=COLORS["primary"])

    chart = BarChart()
    chart.type = "col"
    chart.style = 11
    chart.title = "月別 収入 vs 支出 推移"
    chart.y_axis.title = "金額（円）"
    chart.y_axis.numFmt = '#,##0"円"'
    chart.y_axis.majorGridlines = openpyxl.chart.axis.ChartLines()
    chart.x_axis.title = "月"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    data = Reference(ws, min_col=3, min_row=tbl_hdr_row, max_col=4, max_row=tbl_hdr_row + 12)
    cats = Reference(ws, min_col=2, min_row=tbl_hdr_row + 1, max_row=tbl_hdr_row + 12)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.width = 26
    chart.height = 14
    chart.gapWidth = 80
    chart.overlap = -10
    # Legend position
    chart.legend.position = "t"
    # Color the series
    chart.series[0].graphicalProperties.solidFill = COLORS["success"]  # Income = green
    chart.series[0].graphicalProperties.line.solidFill = COLORS["success"]
    chart.series[1].graphicalProperties.solidFill = COLORS["danger"]   # Expense = red
    chart.series[1].graphicalProperties.line.solidFill = COLORS["danger"]
    # Add data labels
    from openpyxl.chart.label import DataLabelList
    chart.series[0].dLbls = DataLabelList(showVal=False)
    chart.series[1].dLbls = DataLabelList(showVal=False)
    ws.add_chart(chart, f"B{chart_start_row + 1}")

    # Savings trend line chart
    savings_chart_row = chart_start_row + 18
    ws.merge_cells(f"B{savings_chart_row}:D{savings_chart_row}")
    ws[f"B{savings_chart_row}"] = "💰 月別貯蓄額の推移"
    ws[f"B{savings_chart_row}"].font = Font(name="Arial", size=12, bold=True, color=COLORS["primary"])

    line_chart = LineChart()
    line_chart.style = 12
    line_chart.title = "月別貯蓄額の推移"
    line_chart.y_axis.title = "貯蓄額（円）"
    line_chart.y_axis.numFmt = '#,##0"円"'
    line_chart.y_axis.majorGridlines = openpyxl.chart.axis.ChartLines()
    line_chart.x_axis.title = "月"
    line_chart.x_axis.delete = False
    line_chart.y_axis.delete = False
    savings_data = Reference(ws, min_col=5, min_row=tbl_hdr_row, max_row=tbl_hdr_row + 12)
    line_chart.add_data(savings_data, titles_from_data=True)
    line_chart.set_categories(cats)
    line_chart.width = 26
    line_chart.height = 12
    line_chart.legend.position = "t"
    line_chart.series[0].graphicalProperties.line.solidFill = COLORS["accent"]
    line_chart.series[0].graphicalProperties.line.width = 30000  # EMU units (≈2.4pt)
    # Add markers to the line for graphical appeal
    from openpyxl.chart.marker import Marker
    line_chart.series[0].marker = Marker(symbol="circle", size=8)
    line_chart.series[0].marker.graphicalProperties = openpyxl.chart.shapes.GraphicalProperties(solidFill=COLORS["accent"])
    ws.add_chart(line_chart, f"B{savings_chart_row + 1}")

    # ===== Category Breakdown Section (right side) =====
    # Put expense category summary on the right
    cat_col = 10  # Column J
    ws.column_dimensions[get_column_letter(cat_col)].width = 22
    ws.column_dimensions[get_column_letter(cat_col + 1)].width = 14
    ws.column_dimensions[get_column_letter(cat_col + 2)].width = 10

    ws.merge_cells(f"{get_column_letter(cat_col)}{section_row}:{get_column_letter(cat_col+2)}{section_row}")
    ws.cell(row=section_row, column=cat_col, value="🏷️ カテゴリ別 年間支出")
    ws.cell(row=section_row, column=cat_col).font = Font(name="Arial", size=13, bold=True, color=COLORS["primary"])

    cat_hdr_row = section_row + 1
    cat_headers = ["🏷️ カテゴリ", "💴 年間合計", "📊 構成比"]
    for col, h in enumerate(cat_headers, cat_col):
        ws.cell(row=cat_hdr_row, column=col, value=h)
        ws.cell(row=cat_hdr_row, column=col).font = header_font
        ws.cell(row=cat_hdr_row, column=col).fill = header_fill
        ws.cell(row=cat_hdr_row, column=col).alignment = Alignment(horizontal="center")
        ws.cell(row=cat_hdr_row, column=col).border = thin_border

    for i, cat in enumerate(EXPENSE_CATEGORIES):
        r = cat_hdr_row + 1 + i
        ws.cell(row=r, column=cat_col, value=cat).border = thin_border
        # Sum across all months for this category (row offset in monthly sheets)
        # In monthly sheet, expenses start at row 11 (row 10+1)
        expense_row_in_month = 11 + i
        parts = [f"'{m}'!D{expense_row_in_month}" for m in MONTHS]
        ws.cell(row=r, column=cat_col + 1, value=f"={'+'.join(parts)}")
        ws.cell(row=r, column=cat_col + 1).number_format = '#,##0"円"'
        ws.cell(row=r, column=cat_col + 1).border = thin_border
        # Percentage
        total_ref = f"D{val_row}"
        ws.cell(row=r, column=cat_col + 2, value=f'=IF({total_ref}=0,"",{get_column_letter(cat_col+1)}{r}/{total_ref})')
        ws.cell(row=r, column=cat_col + 2).number_format = "0%"
        ws.cell(row=r, column=cat_col + 2).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=cat_col + 2).border = thin_border
        # Alternate rows
        if i % 2 == 0:
            for c in range(cat_col, cat_col + 3):
                ws.cell(row=r, column=c).fill = PatternFill(
                    start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

    # Pie chart for category breakdown
    pie_row = cat_hdr_row + len(EXPENSE_CATEGORIES) + 2
    pie = PieChart()
    pie.title = "支出カテゴリ構成比"
    pie.style = 10
    pie_data = Reference(ws, min_col=cat_col + 1, min_row=cat_hdr_row,
                         max_row=cat_hdr_row + len(EXPENSE_CATEGORIES))
    pie_cats = Reference(ws, min_col=cat_col, min_row=cat_hdr_row + 1,
                         max_row=cat_hdr_row + len(EXPENSE_CATEGORIES))
    pie.add_data(pie_data, titles_from_data=True)
    pie.set_categories(pie_cats)
    pie.width = 16
    pie.height = 12
    ws.add_chart(pie, f"{get_column_letter(cat_col)}{pie_row}")


def create_monthly_sheet(wb, month_name):
    month_idx = MONTHS.index(month_name)
    icon = MONTH_ICONS[month_idx]
    ws = wb.create_sheet(title=month_name)
    set_col_widths(ws, [3, 22, 15, 15, 15, 15, 3])

    # Title
    ws.merge_cells("B1:F1")
    ws[f"B1"] = f"{icon} {month_name} 家計簿"
    ws[f"B1"].font = title_font

    # === Income Section ===
    row = 3
    headers = ["🏷️ カテゴリ", "📋 予算", "✅ 実績", "📊 合計", "📝 メモ"]
    ws.cell(row=row, column=2, value="💚【収入】")
    ws.cell(row=row, column=2).font = subtitle_font
    ws.cell(row=row, column=2).fill = income_fill
    # Income total formula
    income_start = row + 1
    income_end = row + len(INCOME_CATEGORIES)
    ws.cell(row=row, column=5, value=f"=SUM(D{income_start}:D{income_end})")
    ws.cell(row=row, column=5).font = subtitle_font
    ws.cell(row=row, column=5).number_format = '#,##0"円"'

    row = 4
    # Category rows for income
    for i, cat in enumerate(INCOME_CATEGORIES):
        r = row + i
        ws.cell(row=r, column=2, value=cat).border = thin_border
        ws.cell(row=r, column=3, value=0).border = thin_border  # Budget
        ws.cell(row=r, column=3).number_format = '#,##0'
        ws.cell(row=r, column=4, value=0).border = thin_border  # Actual
        ws.cell(row=r, column=4).number_format = '#,##0'
        ws.cell(row=r, column=5, value="").border = thin_border  # Diff (user can add)
        ws.cell(row=r, column=6, value="").border = thin_border  # Memo

    # === Expense Section ===
    row = 4 + len(INCOME_CATEGORIES) + 2
    ws.cell(row=row, column=2, value="🔴【支出】")
    ws.cell(row=row, column=2).font = subtitle_font
    ws.cell(row=row, column=2).fill = expense_fill
    expense_start = row + 1
    expense_end = row + len(EXPENSE_CATEGORIES)
    ws.cell(row=row, column=5, value=f"=SUM(D{expense_start}:D{expense_end})")
    ws.cell(row=row, column=5).font = subtitle_font
    ws.cell(row=row, column=5).number_format = '#,##0"円"'

    row += 1
    for i, cat in enumerate(EXPENSE_CATEGORIES):
        r = row + i
        ws.cell(row=r, column=2, value=cat).border = thin_border
        ws.cell(row=r, column=3, value=0).border = thin_border
        ws.cell(row=r, column=3).number_format = '#,##0'
        ws.cell(row=r, column=4, value=0).border = thin_border
        ws.cell(row=r, column=4).number_format = '#,##0'
        ws.cell(row=r, column=5, value="").border = thin_border
        ws.cell(row=r, column=6, value="").border = thin_border

    # === Balance Section ===
    bal_row = row + len(EXPENSE_CATEGORIES) + 2
    ws.cell(row=bal_row, column=2, value="⚖️【収支バランス】")
    ws.cell(row=bal_row, column=2).font = subtitle_font
    ws.cell(row=bal_row, column=2).fill = savings_fill

    ws.cell(row=bal_row + 1, column=2, value="⬆️ 収入合計").border = thin_border
    ws.cell(row=bal_row + 1, column=4, value=f"=E3").border = thin_border
    ws.cell(row=bal_row + 1, column=4).number_format = '#,##0"円"'

    ws.cell(row=bal_row + 2, column=2, value="⬇️ 支出合計").border = thin_border
    expense_header_row = 4 + len(INCOME_CATEGORIES) + 2
    ws.cell(row=bal_row + 2, column=4, value=f"=E{expense_header_row}").border = thin_border
    ws.cell(row=bal_row + 2, column=4).number_format = '#,##0"円"'

    ws.cell(row=bal_row + 3, column=2, value="✨ 差額（貯蓄）").border = thin_border
    ws.cell(row=bal_row + 3, column=2).font = Font(name="Arial", size=10, bold=True)
    ws.cell(row=bal_row + 3, column=4, value=f"=D{bal_row+1}-D{bal_row+2}").border = thin_border
    ws.cell(row=bal_row + 3, column=4).number_format = '#,##0"円"'

    # === Daily Expense Log ===
    log_row = bal_row + 6
    ws.cell(row=log_row, column=2, value="📝【日別支出記録】")
    ws.cell(row=log_row, column=2).font = subtitle_font

    log_headers = ["📅 日付（何日）", "🏷️ カテゴリ（種類）", "💴 金額（円）", "💳 支払方法", "📝 メモ・備考"]
    log_row += 1
    for col, h in enumerate(log_headers, 2):
        ws.cell(row=log_row, column=col, value=h)
    style_header_row(ws, log_row, 6)

    # Empty rows for daily logging (31 days)
    for i in range(1, 32):
        r = log_row + i
        for c in range(2, 7):
            ws.cell(row=r, column=c).border = thin_border

    # Data validation: Category dropdown for daily expense log
    category_list = ",".join([c.split(" ", 1)[1] if " " in c else c for c in EXPENSE_CATEGORIES])
    dv_category = DataValidation(
        type="list",
        formula1=f'"{category_list}"',
        allow_blank=True,
    )
    dv_category.error = "リストからカテゴリを選択してください"
    dv_category.errorTitle = "カテゴリエラー"
    dv_category.prompt = "カテゴリを選んでください"
    dv_category.promptTitle = "カテゴリ選択"
    dv_category.add(f"C{log_row + 1}:C{log_row + 31}")
    ws.add_data_validation(dv_category)

    # Data validation: Payment method dropdown
    payment_methods = "現金,クレジットカード,デビットカード,電子マネー,QRコード決済,口座振替,その他"
    dv_payment = DataValidation(
        type="list",
        formula1=f'"{payment_methods}"',
        allow_blank=True,
    )
    dv_payment.error = "リストから支払方法を選択してください"
    dv_payment.errorTitle = "支払方法エラー"
    dv_payment.prompt = "支払方法を選んでください"
    dv_payment.promptTitle = "支払方法選択"
    dv_payment.add(f"E{log_row + 1}:E{log_row + 31}")
    ws.add_data_validation(dv_payment)


def create_bills_tracker(wb):
    ws = wb.create_sheet(title="固定費・サブスク管理")
    set_col_widths(ws, [3, 22, 12, 12, 12, 12, 15, 15])

    ws.merge_cells("B1:G1")
    ws["B1"] = "📌 固定費・サブスクリプション管理"
    ws["B1"].font = title_font

    headers = ["📋 項目名", "💴 月額", "📅 支払日", "💳 支払方法", "🏷️ カテゴリ", "⏰ 次回支払日", "✅ ステータス"]
    row = 3
    for col, h in enumerate(headers, 2):
        ws.cell(row=row, column=col, value=h)
    style_header_row(ws, row, 8)

    # Sample entries - 大幅にカテゴリを拡充
    bills = [
        # === 住居費 ===
        ("🏠 家賃 / 住宅ローン", 80000, 25, "口座振替", "住居費", "", "継続中"),
        ("🏢 管理費・共益費", 8000, 25, "口座振替", "住居費", "", "継続中"),
        ("🅿️ 駐車場代", 12000, 27, "口座振替", "住居費", "", "継続中"),
        ("🏘️ 火災・地震保険", 1500, 1, "口座振替", "保険料", "", "継続中"),
        # === 水道光熱費 ===
        ("⚡ 電気代", 8000, 15, "口座振替", "水道光熱費", "", "継続中"),
        ("🔥 ガス代", 5000, 15, "口座振替", "水道光熱費", "", "継続中"),
        ("💧 水道代", 4000, 15, "口座振替", "水道光熱費", "", "継続中"),
        # === 通信費 ===
        ("📱 携帯電話", 5000, 10, "クレジットカード", "通信費", "", "継続中"),
        ("🌐 インターネット回線", 5000, 20, "クレジットカード", "通信費", "", "継続中"),
        ("📡 NHK受信料", 1275, 1, "口座振替", "通信費", "", "継続中"),
        # === 動画・音楽サブスク ===
        ("🎬 Netflix", 1490, 1, "クレジットカード", "娯楽費", "", "継続中"),
        ("🎞️ Amazon Prime", 600, 15, "クレジットカード", "娯楽費", "", "継続中"),
        ("📺 Disney+", 990, 1, "クレジットカード", "娯楽費", "", "継続中"),
        ("🎥 Hulu", 1026, 1, "クレジットカード", "娯楽費", "", "継続中"),
        ("📼 U-NEXT", 2189, 1, "クレジットカード", "娯楽費", "", "停止中"),
        ("🎵 Spotify", 980, 1, "クレジットカード", "娯楽費", "", "継続中"),
        ("🎧 Apple Music", 1080, 1, "クレジットカード", "娯楽費", "", "停止中"),
        ("📻 YouTube Premium", 1280, 1, "クレジットカード", "娯楽費", "", "継続中"),
        # === ソフトウェア / クラウド ===
        ("☁️ iCloud+ ストレージ", 400, 1, "クレジットカード", "通信費", "", "継続中"),
        ("📁 Google One", 250, 1, "クレジットカード", "通信費", "", "継続中"),
        ("💼 Microsoft 365", 1284, 1, "クレジットカード", "通信費", "", "継続中"),
        ("🎨 Adobe Creative Cloud", 6480, 1, "クレジットカード", "教育費", "", "停止中"),
        ("🤖 ChatGPT Plus", 3000, 1, "クレジットカード", "教育費", "", "継続中"),
        ("📔 Notion", 1000, 1, "クレジットカード", "教育費", "", "停止中"),
        # === 保険 ===
        ("🛡️ 生命保険", 10000, 27, "口座振替", "保険料", "", "継続中"),
        ("🏥 医療保険", 4000, 27, "口座振替", "保険料", "", "継続中"),
        ("🚗 自動車保険", 5000, 27, "口座振替", "保険料", "", "継続中"),
        ("👨‍👩‍👧 学資保険", 10000, 27, "口座振替", "保険料", "", "停止中"),
        # === 健康・フィットネス ===
        ("💪 ジム会費", 8000, 5, "クレジットカード", "美容費", "", "継続中"),
        ("🧘 ヨガスタジオ", 12000, 5, "クレジットカード", "美容費", "", "停止中"),
        # === 教育 ===
        ("📚 英会話レッスン", 6000, 1, "クレジットカード", "教育費", "", "停止中"),
        ("📖 新聞購読", 4400, 1, "口座振替", "教育費", "", "停止中"),
        # === 交通 ===
        ("🚃 定期券", 15000, 1, "クレジットカード", "交通費", "", "継続中"),
        # === その他 ===
        ("💳 クレジットカード年会費", 11000, 5, "口座振替", "その他", "", "継続中"),
        ("🎁 ふるさと納税（積立）", 5000, 25, "クレジットカード", "その他", "", "継続中"),
        # === 予備行 ===
        ("", 0, "", "", "", "", ""),
        ("", 0, "", "", "", "", ""),
        ("", 0, "", "", "", "", ""),
    ]

    for i, bill in enumerate(bills):
        r = row + 1 + i
        for c, val in enumerate(bill, 2):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = thin_border
            if c == 3 and val:  # Amount
                cell.number_format = '#,##0"円"'

    # Total
    total_row = row + 1 + len(bills) + 1
    ws.cell(row=total_row, column=2, value="📊 月額合計").font = subtitle_font
    ws.cell(row=total_row, column=3, value=f"=SUM(C4:C{row+len(bills)})").number_format = '#,##0"円"'
    ws.cell(row=total_row, column=3).font = subtitle_font


def create_savings_tracker(wb):
    ws = wb.create_sheet(title="貯蓄目標")
    set_col_widths(ws, [3, 22, 15, 15, 15, 12, 15])

    ws.merge_cells("B1:F1")
    ws["B1"] = "🎯 貯蓄目標トラッカー"
    ws["B1"].font = title_font

    headers = ["🎯 目標名", "💰 目標金額", "📊 現在の貯蓄", "📉 残り", "🏆 達成率", "📅 期限"]
    row = 3
    for col, h in enumerate(headers, 2):
        ws.cell(row=row, column=col, value=h)
    style_header_row(ws, row, 7)

    goals = [
        ("🆘 緊急予備費", 1000000, 300000),
        ("✈️ 旅行資金", 500000, 150000),
        ("🚗 車購入", 2000000, 500000),
        ("🎓 教育資金", 3000000, 800000),
        ("🏖️ 老後資金", 10000000, 1000000),
    ]

    for i, (name, target, current) in enumerate(goals):
        r = row + 1 + i
        ws.cell(row=r, column=2, value=name).border = thin_border
        ws.cell(row=r, column=3, value=target).border = thin_border
        ws.cell(row=r, column=3).number_format = '#,##0"円"'
        ws.cell(row=r, column=4, value=current).border = thin_border
        ws.cell(row=r, column=4).number_format = '#,##0"円"'
        ws.cell(row=r, column=5, value=f"=C{r}-D{r}").border = thin_border
        ws.cell(row=r, column=5).number_format = '#,##0"円"'
        ws.cell(row=r, column=6, value=f"=D{r}/C{r}").border = thin_border
        ws.cell(row=r, column=6).number_format = "0%"
        ws.cell(row=r, column=7, value="").border = thin_border

    # Add data bars for progress
    rule = DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                       color="27AE60")
    ws.conditional_formatting.add(f"F4:F{row+len(goals)}", rule)


def create_instructions(wb):
    ws = wb.create_sheet(title="使い方ガイド")
    set_col_widths(ws, [3, 80])

    ws["B1"] = "📖 家計簿の使い方ガイド"
    ws["B1"].font = title_font

    instructions = [
        "",
        "📌【はじめに】",
        "この家計簿テンプレートは、毎月の収入と支出を管理し、",
        "年間の貯蓄目標を達成するためのツールです。",
        "",
        "🔰【基本的な使い方】",
        "1️⃣ 各月のシートに毎日の収入と支出を記録します",
        "2️⃣ カテゴリごとの予算を設定します（予算列に入力）",
        "3️⃣ 実績を入力すると、自動的に合計が計算されます",
        "4️⃣ ダッシュボードで年間のサマリーを確認できます",
        "",
        "📑【シートの説明】",
        "• 📊 ダッシュボード：年間の収支サマリーとグラフ",
        "• 📅 1月〜12月：各月の収支管理と日別支出記録",
        "• 📌 固定費・サブスク管理：毎月の固定費一覧",
        "• 🎯 貯蓄目標：貯蓄目標の進捗管理",
        "",
        "🔧【カスタマイズ方法】",
        "• ✏️ カテゴリは自由に追加・変更できます",
        "• 🎨 色やフォントはお好みで変更してください",
        "• ➕ 行を追加する場合は、数式が反映されるか確認してください",
        "",
        "💡【Tips】",
        "• ✍️ 毎日記録する習慣をつけましょう",
        "• 🗓️ 月初に予算を設定し、月末に振り返りをしましょう",
        "• 🔄 固定費は「固定費・サブスク管理」シートで一括管理",
        "• 🚀 貯蓄目標を設定するとモチベーションが上がります",
        "",
        "⚠️【注意事項】",
        "• 🌐 このファイルはGoogleスプレッドシートにインポートして使えます",
        "• 🔒 数式のあるセルを上書きしないよう注意してください",
        "• 💾 定期的にバックアップを取ることをおすすめします",
    ]

    for i, line in enumerate(instructions, 2):
        ws.cell(row=i, column=2, value=line)
        if "【" in line:
            ws.cell(row=i, column=2).font = subtitle_font


def add_sample_data(wb):
    """Add sample data to first 3 months so the dashboard looks good."""
    import random
    random.seed(42)

    # Sample monthly income (salary + side income)
    monthly_incomes = [
        [300000, 50000, 5000, 0],   # Jan
        [300000, 45000, 5000, 10000],  # Feb
        [300000, 60000, 8000, 0],   # Mar
    ]

    # Sample monthly expenses per category
    monthly_expenses = [
        [80000, 15000, 45000, 8000, 12000, 10000, 15000, 3000, 0, 15000, 5000, 3000, 10000, 5000],
        [80000, 13000, 42000, 7000, 11000, 10000, 15000, 5000, 0, 12000, 3000, 4000, 8000, 4000],
        [80000, 12000, 48000, 9000, 13000, 10000, 15000, 0, 5000, 20000, 8000, 3000, 15000, 6000],
    ]

    for month_idx in range(3):
        month_name = MONTHS[month_idx]
        ws = wb[month_name]

        # Fill income (rows 4-7, column D = actual)
        for i, amount in enumerate(monthly_incomes[month_idx]):
            ws.cell(row=4 + i, column=4, value=amount)

        # Fill expenses (rows 11-24, column D = actual)
        for i, amount in enumerate(monthly_expenses[month_idx]):
            ws.cell(row=11 + i, column=4, value=amount)


def main():
    wb = Workbook()

    # Create all sheets
    create_dashboard(wb)
    for month in MONTHS:
        create_monthly_sheet(wb, month)
    create_bills_tracker(wb)
    create_savings_tracker(wb)
    create_instructions(wb)

    # Add sample data for visual demonstration
    add_sample_data(wb)

    wb.save(OUTPUT_PATH)
    print(f"✅ 家計簿テンプレートを作成しました: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
